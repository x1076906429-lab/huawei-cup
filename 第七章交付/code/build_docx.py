# -*- coding: utf-8 -*-
"""将 问题三第七章正文.md 转为 Word：公式为可编辑的 Word 公式（OMML），图片内嵌。
排版：A4，页边距 2.5 cm；正文宋体 + Times New Roman 小四、1.5 倍行距、首行缩进 2 字符；
标题黑体（章 三号居中、节 四号、小节 小四）；图题/表题 五号居中（表题在上、图题在下）；
三线表；公式居中、编号右对齐。用法：python build_docx.py"""
import os, re, copy
import pypandoc
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

D = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
SRC = os.path.join(D, '问题三第七章正文.md')
OUT = os.path.join(D, '问题三第七章正文.docx')
CN, CN_H, EN = '宋体', '黑体', 'Times New Roman'

# ---------- 1. 预处理 Markdown ----------
md = open(SRC, encoding='utf-8').read()
md = re.sub(r'!\[[^\]]*\]\((figures/[^)]+)\)', lambda m: f'![]({m.group(1)}){{width=15cm}}', md)   # 空 alt：不生成重复图题
md = re.sub(r'\$\$\s*(.*?)\s*\\tag\{(7-\d+)\}\s*\$\$', lambda m: f'$$\n{m.group(1)}\n$$\n\nEQNUM({m.group(2)})', md, flags=re.S)
tmp = os.path.join(D, '_tmp_build.md')
open(tmp, 'w', encoding='utf-8').write(md)
pypandoc.convert_file(tmp, 'docx', format='markdown+tex_math_dollars+pipe_tables', outputfile=OUT,
                      extra_args=[f'--resource-path={D}'])
os.remove(tmp)

# ---------- 2. 后处理 ----------
doc = Document(OUT)


def set_font(run_or_style, cn=CN, en=EN, size=None, bold=None):
    f = run_or_style.font
    f.name = en
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    f.color.rgb = RGBColor(0, 0, 0)
    el = run_or_style.element if hasattr(run_or_style, 'element') else run_or_style._element
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs'): rf.set(qn(k), en)
    rf.set(qn('w:eastAsia'), cn)
    for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        rf.attrib.pop(qn(k), None)


def pfmt(p, align=None, indent=True, before=0, after=0, spacing=1.5):
    pf = p.paragraph_format
    if align is not None: pf.alignment = align
    pf.first_line_indent = Pt(24) if indent else Pt(0)
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = spacing


# 页面
for s in doc.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)
    # 页脚页码
    fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    for t, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if t:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t); r._r.append(e)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._r.append(e)
    set_font(r, size=10.5)

# 样式
for name in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
    if name in [s.name for s in doc.styles]:
        set_font(doc.styles[name], size=12)
HEAD = {'Heading 1': (16, WD_ALIGN_PARAGRAPH.CENTER, 12, 12), 'Heading 2': (14, WD_ALIGN_PARAGRAPH.LEFT, 12, 6),
        'Heading 3': (12, WD_ALIGN_PARAGRAPH.LEFT, 6, 6)}
for name, (sz, al, b, a) in HEAD.items():
    st = doc.styles[name]; set_font(st, cn=CN_H, size=sz, bold=True)
    st.font.italic = False
    st.paragraph_format.alignment = al
    st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(b), Pt(a)
    st.paragraph_format.first_line_indent = Pt(0)
    st.paragraph_format.line_spacing = 1.5

body = doc.element.body


def text_of(p):
    return ''.join(t.text or '' for t in p._p.iter(qn('w:t')))


def has_math(p):
    return p._p.find('.//' + qn('m:oMathPara')) is not None


def has_img(p):
    return p._p.find('.//' + qn('w:drawing')) is not None


def no_border_table(ncols, widths):
    t = doc.add_table(rows=1, cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'nil'); b.append(x)
    tblPr.append(b)
    for c, w in zip(t.rows[0].cells, widths): c.width = Cm(w)
    return t


# ---------- 公式：1×3 无框表，公式居中、编号右对齐 ----------
paras = list(doc.paragraphs)
for i, p in enumerate(paras):
    m = re.fullmatch(r'EQNUM\((7-\d+)\)', text_of(p).strip())
    if not m: continue
    eqp = paras[i - 1]
    assert has_math(eqp), m.group(1)
    t = no_border_table(3, [1.5, 13, 1.5])
    c0, c1, c2 = t.rows[0].cells
    c1.paragraphs[0]._p.append(copy.deepcopy(eqp._p.find(qn('m:oMathPara'))))
    c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    rp = c2.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(rp.add_run(f'({m.group(1)})'), size=12)
    for c in (c0, c1, c2):
        c.vertical_alignment = 1
        pfmt(c.paragraphs[0], indent=False, spacing=1.0, before=4, after=4)
    c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    eqp._p.addprevious(t._tbl)
    eqp._p.getparent().remove(eqp._p); p._p.getparent().remove(p._p)

# ---------- 段落格式 ----------
for p in doc.paragraphs:
    txt = text_of(p).strip()
    sname = p.style.name
    if sname.startswith('Heading'):
        for r in p.runs: set_font(r, cn=CN_H, size=HEAD[sname][0], bold=True)
        continue
    if has_img(p):
        pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=6, spacing=1.0)
        p.paragraph_format.keep_with_next = True   # 图与图题同页
        continue
    if re.match(r'^[图表]7-\d+ ', txt):          # 图题 / 表题
        is_tab = txt.startswith('表')
        pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=6 if is_tab else 0, after=3 if is_tab else 9, spacing=1.25)
        for r in p.runs: set_font(r, cn=CN_H if r.bold else CN, size=10.5)
        if is_tab: p.paragraph_format.keep_with_next = True   # 表题与表同页
        continue
    if txt.startswith('注：') or txt.startswith('（编号为本章临时编号'):
        pfmt(p, indent=False, after=6, spacing=1.25)
        for r in p.runs: set_font(r, size=10.5)
        continue
    if re.match(r'^\[\d\] ', txt):              # 参考文献
        pf = p.paragraph_format; pfmt(p, indent=False, spacing=1.25)
        pf.left_indent, pf.first_line_indent = Pt(21), Pt(-21)
        for r in p.runs: set_font(r, size=10.5)
        continue
    pfmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True)
    for r in p.runs:
        code = r.style is not None and r.style.name == 'Verbatim Char'
        set_font(r, en=EN, size=12)
        if code: r.font.name = EN  # 代码体统一为 Times New Roman，不用等宽字体

# ---------- 三线表 ----------
def border(el, edge, sz):
    x = OxmlElement(f'w:{edge}')
    x.set(qn('w:val'), 'single' if sz else 'nil'); x.set(qn('w:sz'), str(sz)); x.set(qn('w:space'), '0'); x.set(qn('w:color'), '000000')
    el.append(x)


for t in doc.tables:
    first = t.rows[0].cells[0]
    if len(t.columns) == 3 and first.paragraphs[0].text == '' and len(t.rows) == 1:
        continue  # 公式表
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')): tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for e, sz in (('top', 12), ('left', 0), ('bottom', 12), ('right', 0), ('insideH', 0), ('insideV', 0)): border(b, e, sz)
    tblPr.append(b)
    tw = tblPr.find(qn('w:tblW'))
    if tw is None: tw = OxmlElement('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), str(int(16 * 567)))
    # 列宽：按表体最长文本的实际字符宽度分配（中文计 2 个半角单位），表头允许折两行；
    # 总宽超过版心时改用 9 pt 小字号，最后按比例缩放到 16 cm；固定布局，单元格左右边距 0.1 cm
    import math
    def vlen(s): return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)
    ncol = len(t.columns)
    need = []
    for c in range(ncol):
        bodyw = max([min(vlen(t.rows[r].cells[c].text), 44) for r in range(1, len(t.rows))] or [0])
        need.append(max(bodyw, math.ceil(vlen(t.rows[0].cells[c].text) / 2), 4))
    fs = 10.5
    unit = lambda f: 0.0353 * f * 0.52           # 每个半角单位的宽度（cm）
    raw = [n * unit(fs) + 0.45 for n in need]
    if sum(raw) > 16.0:
        fs = 9.0; raw = [n * unit(fs) + 0.35 for n in need]
    cw = [w * 16.0 / sum(raw) for w in raw]
    t._fs = fs
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    mar = OxmlElement('w:tblCellMar')
    for e in ('left', 'right'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:w'), '57'); x.set(qn('w:type'), 'dxa'); mar.append(x)
    tblPr.append(mar)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), cw): gc.set(qn('w:w'), str(int(w * 567)))
    for row in t.rows:
        for c, w in zip(row.cells, cw): c.width = Cm(w)
    heads = [c.text for c in t.rows[0].cells]
    left_cols = {i for i, h in enumerate(heads) if h in ('取值', '样本', '样本分组', '项目', '指标', '检验')}
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            if ri == 0:
                cb = OxmlElement('w:tcBorders'); border(cb, 'bottom', 6); tcPr.append(cb)
            cell.vertical_alignment = 1
            for p in cell.paragraphs:
                pfmt(p, align=WD_ALIGN_PARAGRAPH.LEFT if ci in left_cols and ri > 0 else WD_ALIGN_PARAGRAPH.CENTER,
                     indent=False, spacing=1.0, before=1.5, after=1.5)
                for r in p.runs: set_font(r, size=t._fs, bold=(ri == 0))
    # 表头重复
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)

doc.save(OUT)
print('saved', OUT)
